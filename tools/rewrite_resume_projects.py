from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SOURCE = Path(r"C:\Users\艾米\Desktop\简历_双AI项目版_项目内容优化_原模板.docx")
OUTPUT = Path(r"C:\Users\艾米\Desktop\简历_双AI项目版_项目内容优化_参考图风格版.docx")

BLUE = RGBColor(31, 78, 121)
LINK_BLUE = RGBColor(5, 99, 193)
FONT = "Microsoft YaHei"


def clear_paragraph(paragraph):
    """Remove all content while preserving the paragraph properties."""
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_run_font(run, size=6.5, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def ensure_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    if numbering.xpath('./w:num[@w:numId="2"]'):
        return 2

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "1")
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "360")
    ind.set(qn("w:hanging"), "180")
    ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Symbol")
    rfonts.set(qn("w:hAnsi"), "Symbol")
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "2")
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), "1")
    num.append(abstract_id)
    numbering.append(num)
    return 2


def apply_bullet(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:numPr"))
    if old is not None:
        ppr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(num_id_el)
    ppr.append(num_pr)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "13")
    rpr.append(size)
    run.append(rpr)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_title(paragraph, title, stack):
    clear_paragraph(paragraph)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, size=8.0, bold=True, color=BLUE)
    sep = paragraph.add_run("  |  ")
    set_run_font(sep, size=6.5, color=BLUE)
    stack_run = paragraph.add_run(stack)
    set_run_font(stack_run, size=6.5, color=BLUE)
    paragraph.paragraph_format.keep_with_next = True


def add_summary(paragraph, text):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=6.5)


def add_bullet(paragraph, label, text, num_id, link=None):
    clear_paragraph(paragraph)
    apply_bullet(paragraph, num_id)
    label_run = paragraph.add_run(label + ": ")
    set_run_font(label_run, size=6.5, bold=True, color=BLUE)
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size=6.5)
    if link:
        prefix = paragraph.add_run("  ")
        set_run_font(prefix, size=6.5)
        add_hyperlink(paragraph, "GitHub 源码", link)


def main():
    doc = Document(SOURCE)
    paragraphs = doc.paragraphs
    bullet_num_id = ensure_bullet_numbering(doc)

    add_title(
        paragraphs[12],
        "QAU SmartQA · 校园智能问答平台（开源项目）",
        "Java / Spring AI / pgvector / Redis / Vue / SSE",
    )
    add_summary(
        paragraphs[13],
        "面向校园官网老旧、信息分散，学生和教职工查询政策、流程、通知成本高的问题，独立构建统一 AI 问答入口；围绕校园政策、办事指南、通知公告等资料，组织多知识库选择、文档解析切片、向量检索、会话记忆与回答来源追踪，形成从资料入库、检索问答到历史复盘的校园智能助手基础。",
    )
    add_bullet(
        paragraphs[14],
        "Context Engineering",
        "针对校园政策、流程和通知需要准确引用的问题，设计 Prompt、上下文组装、历史消息裁剪与引用来源分离机制，让模型回答范围可控、来源可追溯。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[15],
        "Hybrid RAG",
        "围绕官网资料后续接入需求，设计文档、切片、Embedding、pgvector TopK 检索与重排链路，并预留 Tool Calling 接入官网数据库和校内资料，支持从检索式问答扩展到可操作助手。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[16],
        "Agent Runtime",
        "基于 Spring AI ChatClient 封装统一模型调用、SSE 事件和异常降级，按 qau-app、qau-ai、qau-domain、qau-infrastructure、qau-common 拆分模块，保持 AI 能力与业务领域解耦。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[17],
        "可靠性与可观测性",
        "针对模型超时、客户端断开、上下文过长导致的请求悬挂和回答失控，补充超时/降级、异常事件、会话状态清理和流式输出边界，提升 Agent 行为的可恢复性。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[18],
        "工程化落地",
        "完成登录鉴权、会话创建、历史消息加载、MyBatis-Plus CRUD、Flyway 建表和 Vue 联调；以 PostgreSQL + pgvector、Redis、SSE、Docker 形成可复现的开发与部署基础。",
        bullet_num_id,
        link="https://github.com/aimiaa/QAU-SmartQA",
    )
    # Keep the original paragraph slot as a compact spacer before the next project.
    clear_paragraph(paragraphs[19])
    paragraphs[19].paragraph_format.space_after = Pt(0)

    add_title(
        paragraphs[20],
        "AI Interviewer · 智能AI面试官平台",
        "Java / Spring AI / pgvector / React / WebSocket",
    )
    add_summary(
        paragraphs[21],
        "面向求职者和招聘场景，构建覆盖简历解析、JD 理解、知识库题库、面试安排、文字/语音模拟面试和多模型配置的 AI 面试平台；根据候选人经历、目标岗位和历史表现动态组织面试上下文，通过 Skill 驱动出题、RAG 检索、实时追问、异步评估与报告汇总，形成从资料准备、面试交互到能力反馈的完整闭环。",
    )
    add_bullet(
        paragraphs[22],
        "Multi-provider & Ingestion",
        "使用 LLM Provider 屏蔽 DashScope、DeepSeek、Kimi、GLM、LM Studio 的调用差异，结合 Apache Tika 统一解析 PDF、DOCX、Markdown，为不同模型和资料来源提供稳定输入。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[23],
        "Skill-driven Planning",
        "按简历、JD、难度、历史题目和 SKILL.md 组装 Prompt，以结构化 JSON 输出主问题、追问和评分 Rubric，使不同岗位与面试模式具备可配置的出题边界。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[24],
        "Real-time Voice Loop",
        "通过 WebSocket 编排 ASR → LLM → TTS，支持实时字幕、暂停/恢复和手动提交；针对 TTS 阻塞与音频乱序，采用句子级并发处理并按序回传，保持交互节奏稳定。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[25],
        "Async Evaluation",
        "使用 Redis Stream 承载简历分析、向量化、出题和评估等长任务，引入任务状态机、幂等记录、失败重试与进度回传，解决重复消费、重复写入和状态不同步问题。",
        bullet_num_id,
    )
    add_bullet(
        paragraphs[26],
        "面试闭环",
        "完成简历解析、知识库入库、文字/语音模拟面试、面试安排、多模型配置和报告导出；对不同面试模式进行分批评分与二次汇总，形成可复用的统一评估引擎。",
        bullet_num_id,
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
